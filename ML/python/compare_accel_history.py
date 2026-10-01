"""Compare causal 1/3/5-second acceleration histories; preserve default model.

Run: python ML/python/compare_accel_history.py
History restarts at each run_id. Requires ordered, uniformly sampled 10 Hz runs.
"""
from copy import copy
import hashlib
import json
from pathlib import Path
from time import perf_counter

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

import train_total_noise_model as training

ROOT = Path(__file__).resolve().parents[1]
MODEL_PATH = ROOT / "python/models/total_noise_model_accel_history_complete.joblib"
TEST_PATH = ROOT / "dataset/GPT_ML_data/RoadSurf_Synthetic_TestSet_02.xlsx"
OUTPUT = ROOT / "dataset/predict_noise_snow/RoadSurf_Synthetic_TestSet_02_predicted_accel_history_complete.xlsx"
REPORT = OUTPUT.with_name("accel_history_complete_comparison.xlsx")
AXES = ["accel_x_g", "accel_y_g", "accel_z_g"]
BASE_FEATURES = training.FEATURE_COLUMNS.copy()


def make_features(data: pd.DataFrame, seconds: int) -> pd.DataFrame:
    """Use current and past samples only; keep the original row index/order."""
    if data.run_id.isna().any():
        raise ValueError("Missing run_id")
    if not data.index.is_unique:
        raise ValueError("Duplicate row index")
    x = data[BASE_FEATURES].apply(pd.to_numeric, errors="coerce").replace([np.inf, -np.inf], np.nan)
    parts = []
    for _, group in data.groupby("run_id", sort=False):
        time = pd.to_numeric(group.time_s, errors="raise").to_numpy(dtype=float)
        if not np.isfinite(time).all() or not np.allclose(np.diff(time), .1, atol=1e-6):
            raise ValueError("Each run must be ordered, uninterrupted 10 Hz data")
        if "split" in group and group.split.nunique(dropna=False) != 1:
            raise ValueError("A run spans multiple splits")
        part = x.loc[group.index].copy()
        n = seconds*10
        if n:
            for axis in AXES:
                values = part[axis]
                rolling = values.rolling(n, min_periods=1)
                part[f"{axis}_mean"] = rolling.mean()
                part[f"{axis}_std"] = rolling.std(ddof=0)
                part[f"{axis}_range"] = rolling.max()-rolling.min()
                part[f"{axis}_rms"] = (values.pow(2).rolling(n, min_periods=1).mean()).pow(.5)
                part[f"{axis}_delta"] = values.diff()
                for lag in (1, 2, 5, 10, 20, 30, 40, 50):
                    if lag <= n:
                        part[f"{axis}_lag_{lag}"] = values.shift(lag)
            # Explicitly inform the model when history is still incomplete.
            part["history_samples_available"] = np.minimum(np.arange(len(group)), n)
        parts.append(part)
    return pd.concat(parts).reindex(data.index)


def score(y, prediction):
    error = np.asarray(prediction)-np.asarray(y)
    if not len(error) or not np.isfinite(error).all():
        raise ValueError("Evaluation requires finite predictions on nonempty rows")
    return {"rows": len(error), "mae_mm": float(np.mean(abs(error))*1000),
            "rmse_mm": float(np.sqrt(np.mean(error**2))*1000),
            "p95_mm": float(np.quantile(abs(error), .95)*1000),
            "max_mm": float(np.max(abs(error))*1000),
            "bias_mm": float(np.mean(error)*1000)}


def check_causality(data):
    sample = data.loc[data.run_id.eq(data.run_id.iloc[0])].iloc[:70].copy()
    full = make_features(sample, 5)
    pd.testing.assert_frame_equal(full.iloc[:35], make_features(sample.iloc[:35], 5))
    altered = sample.copy()
    altered.loc[altered.index[35:], AXES] += 100
    pd.testing.assert_frame_equal(full.iloc[:35], make_features(altered, 5).iloc[:35])
    other = sample.copy()
    other.run_id = "independent_run"
    other.index = np.arange(len(sample), len(sample)*2, dtype=np.int64)
    combined = make_features(pd.concat([sample, other]), 5)
    pd.testing.assert_frame_equal(combined.loc[other.index], make_features(other, 5))
    assert combined.loc[other.index[0], "history_samples_available"] == 0
    assert np.isnan(combined.loc[other.index[0], "accel_x_g_lag_1"])


def write_predictions(source, sheet_name, prediction):
    book = load_workbook(TEST_PATH)
    try:
        sheet = book[sheet_name]
        distance = source.distance_measured_m.to_numpy()
        snow = 3.1715-(distance-prediction)
        names = ["predicted_total_noise_m", "predicted_snow_height_m",
                 "total_noise_absolute_error_m", "snow_height_absolute_error_m"]
        arrays = [prediction, snow, abs(prediction-source.target_total_noise_m.to_numpy()),
                  abs(snow-source.target_snow_depth_m.to_numpy())]
        last = sheet.max_column
        for j, name in enumerate(names, 1):
            cell = sheet.cell(1, last+j, name)
            cell._style = copy(sheet.cell(1, last)._style)
            cell.fill = PatternFill(fill_type="solid", fgColor="FF008577")
            font = copy(cell.font)
            font.color = "FFFFFFFF"
            font.bold = True
            cell.font = font
            sheet.column_dimensions[cell.column_letter].width = 27
        for i, values in enumerate(zip(*arrays), 2):
            for j, value in enumerate(values, 1):
                sheet.cell(i, last+j, float(value) if np.isfinite(value) else None).number_format = "0.000000000"
        book.save(OUTPUT)
    finally:
        book.close()


def main():
    default_hash = hashlib.sha256(training.DEFAULT_MODEL.read_bytes()).hexdigest()
    baseline = joblib.load(training.DEFAULT_MODEL)
    data = pd.read_excel(training.DEFAULT_DATASET, sheet_name=baseline["source_sheet"])
    if not np.isfinite(data[training.TARGET_COLUMN].to_numpy(dtype=float)).all():
        raise ValueError("Non-finite targets; do not remove rows before computing history")
    check_causality(data)
    available = data.groupby("run_id", sort=False).cumcount()
    common_ready = available >= 50
    valid_mask = data.split.eq("validation") & common_ready
    features = {s: make_features(data, s) for s in (0, 1, 3, 5)}
    models = {"baseline_RandomForest": (baseline["model"], 0)}
    records = []
    for seconds in (1, 3, 5):
        x = features[seconds]
        train_mask = data.split.eq("train") & (available >= seconds*10)
        history_columns = [c for c in x if c not in BASE_FEATURES]
        if x.loc[train_mask, history_columns].isna().any().any():
            raise ValueError("Training history is incomplete or contains missing sensor values")
        for name in ("RandomForest", "ExtraTrees", "HistGradientBoosting"):
            candidate = training.build_candidates()[name]
            start = perf_counter()
            candidate.fit(x.loc[train_mask], data.loc[train_mask, training.TARGET_COLUMN])
            elapsed = perf_counter()-start
            key = f"{name}_{seconds}s"
            models[key] = (candidate, seconds)
            result = score(data.loc[valid_mask, training.TARGET_COLUMN], candidate.predict(x.loc[valid_mask]))
            records.append({"candidate": key, "history_seconds": seconds,
                            "feature_count": x.shape[1], "fit_seconds": elapsed,
                            "training_rows": int(train_mask.sum()),
                            **{f"validation_{k}": v for k, v in result.items()}})
            print(f"{key}: validation MAE={result['mae_mm']:.3f} mm ({elapsed:.1f}s)", flush=True)
    best = min(records, key=lambda r: r["validation_mae_mm"])
    winner = best["candidate"]
    model, seconds = models[winner]
    baseline_validation = score(data.loc[valid_mask, training.TARGET_COLUMN],
                                baseline["model"].predict(features[0].loc[valid_mask]))
    records.append({"candidate": "baseline_RandomForest", "history_seconds": 0,
                    "feature_count": len(BASE_FEATURES), "fit_seconds": None,
                    **{f"validation_{k}": v for k, v in baseline_validation.items()}})
    # Select only on validation before opening the external test workbook.
    print(f"Best history model: {winner}", flush=True)
    artifact = {"model": model, "model_type": winner, "feature_columns": list(features[seconds].columns),
                "base_feature_columns": BASE_FEATURES, "history_seconds": seconds,
                "history_recipe": "make_features in compare_accel_history.py; 10 Hz; past/current only; reset per run",
                "target_column": training.TARGET_COLUMN, "source_sheet": baseline["source_sheet"],
                "selection": "minimum validation MAE on common rows at least 5s into each run",
                "warmup_policy": "train only after full window; inference before full window remains blank",
                "beats_baseline_validation": best["validation_mae_mm"] < baseline_validation["mae_mm"]}
    joblib.dump(artifact, MODEL_PATH)
    test = pd.read_excel(TEST_PATH, sheet_name=baseline["source_sheet"])
    test_features = {s: make_features(test, s) for s in (0, 1, 3, 5)}
    test_available = test.groupby("run_id", sort=False).cumcount()
    test_common = test_available >= 50
    evaluations, scenarios, sessions = [], [], []
    for row in records:
        name = row["candidate"]
        candidate, s = models[name]
        for split in ("train", "validation", "test", "ood_test"):
            mask = data.split.eq(split) & common_ready
            result = score(data.loc[mask, training.TARGET_COLUMN], candidate.predict(features[s].loc[mask]))
            evaluations.append({"candidate": name, "dataset": split, **result})
        pred = np.full(len(test), np.nan)
        ready = test_available >= s*10
        pred[ready] = candidate.predict(test_features[s].loc[ready])
        evaluations.append({"candidate": name, "dataset": "TestSet_02",
                            **score(test.loc[test_common, training.TARGET_COLUMN], pred[test_common])})
        if name in (winner, "baseline_RandomForest"):
            for scenario, g in test.loc[test_common].groupby("scenario"):
                scenarios.append({"candidate": name, "scenario": scenario,
                                  **score(g[training.TARGET_COLUMN], pred[g.index])})
            for run, g in test.loc[test_common].groupby("run_id", sort=False):
                sessions.append({"candidate": name, "run_id": run, "scenario": g.scenario.iloc[0],
                                 **score(g[training.TARGET_COLUMN], pred[g.index])})
            evaluations.append({"candidate": name, "dataset": "TestSet_02_own_ready_rows",
                                **score(test.loc[ready, training.TARGET_COLUMN], pred[ready])})
    saved = joblib.load(MODEL_PATH)
    output_ready = test_available >= seconds*10
    predicted = np.full(len(test), np.nan)
    predicted[output_ready] = saved["model"].predict(test_features[seconds].loc[output_ready, saved["feature_columns"]])
    write_predictions(test, baseline["source_sheet"], predicted)
    result = pd.read_excel(OUTPUT, sheet_name=baseline["source_sheet"])
    assert result.loc[~output_ready, result.columns[-4:]].isna().all().all()
    assert result.loc[output_ready, result.columns[-4:]].notna().all().all()
    np.testing.assert_allclose(result.predicted_total_noise_m, predicted, atol=1e-12)
    np.testing.assert_allclose(result.predicted_snow_height_m,
                               3.1715-(test.distance_measured_m-predicted), atol=1e-12)
    np.testing.assert_allclose(result.snow_height_absolute_error_m,
                               abs(result.predicted_snow_height_m-test.target_snow_depth_m), atol=1e-12)
    with pd.ExcelWriter(REPORT, engine="openpyxl") as writer:
        pd.DataFrame(records).sort_values("validation_mae_mm").to_excel(writer, sheet_name="selection", index=False)
        pd.DataFrame(evaluations).to_excel(writer, sheet_name="metrics", index=False)
        pd.DataFrame(scenarios).to_excel(writer, sheet_name="scenarios", index=False)
        pd.DataFrame(sessions).to_excel(writer, sheet_name="sessions", index=False)
        pd.DataFrame({"note": [artifact["history_recipe"], artifact["selection"],
            artifact["warmup_policy"],
            "All model comparisons use common rows >=5s per run; own_ready_rows metrics use each model's window.",
            "TestSet 02 was previously used for error analysis; it is exploratory, not an untouched final holdout.",
            "Default model preserved. History model requires make_features at inference; not a drop-in raw-feature model."]}).to_excel(writer, sheet_name="notes", index=False)
    metadata = {k: v for k, v in artifact.items() if k != "model"}
    metadata["comparison"] = evaluations
    MODEL_PATH.with_suffix(".metrics.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")
    assert default_hash == hashlib.sha256(training.DEFAULT_MODEL.read_bytes()).hexdigest()
    chosen = pd.DataFrame(evaluations)
    print(chosen.loc[chosen.candidate.isin([winner, "baseline_RandomForest"])].round(3).to_string(index=False))
    print("Report:", REPORT)
    print("Predictions:", OUTPUT)
    print("Causality, session reset, saved predictions and unchanged baseline verified.")


if __name__ == "__main__":
    main()
