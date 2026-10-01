"""Append noise/snow predictions and absolute errors against ground truth in metres."""

from copy import copy
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ML_DIR = Path(__file__).resolve().parents[1]
SOURCE = ML_DIR / "dataset/GPT_ML_data/RoadSurf_Synthetic_Labeled_v1_resaved.xlsx"
MODEL = ML_DIR / "python/models/total_noise_model.joblib"
OUTPUT = ML_DIR / "dataset/predict_noise_snow/RoadSurf_Synthetic_Labeled_v1_predicted.xlsx"
NEW_COLUMNS = [
    "predicted_total_noise_m",
    "predicted_snow_height_m",
    "total_noise_absolute_error_m",
    "snow_height_absolute_error_m",
]
TARGET_COLUMNS = ["target_total_noise_m", "target_snow_depth_m"]
REFERENCE_DISTANCE_M = 3.171500
PREDICTION_HEADER_COLOR = "FF008577"  # Teal, distinct from navy/brown source headers.


def main() -> None:
    artifact = joblib.load(MODEL)
    features = artifact["feature_columns"]
    workbook = load_workbook(SOURCE)
    try:
        candidates = [
            sheet for sheet in workbook
            if set(features).issubset(cell.value for cell in sheet[1])
        ]
        if len(candidates) != 1:
            raise ValueError("Expected exactly one sheet containing all model features")
        sheet = candidates[0]
        headers = [cell.value for cell in sheet[1]]
        if set(NEW_COLUMNS).intersection(headers):
            raise ValueError("Source already contains prediction columns")
        # Read model inputs only; ground truth columns never enter predict().
        positions = [headers.index(name) for name in features]
        x = pd.DataFrame(
            ([row[index] for index in positions]
             for row in sheet.iter_rows(min_row=2, values_only=True)),
            columns=features,
        )
        x = x.apply(pd.to_numeric, errors="coerce").replace([np.inf, -np.inf], np.nan)
        if x.empty or x["distance_measured_m"].isna().any():
            raise ValueError("Every data row must have a finite measured distance")
        # Other missing features are handled by the saved pipeline's imputer.
        noise = artifact["model"].predict(x[features])
        snow = REFERENCE_DISTANCE_M - (x["distance_measured_m"].to_numpy() - noise)
        if not np.isfinite(noise).all() or not np.isfinite(snow).all():
            raise ValueError("Non-finite model prediction")

        # Ground truth is used only for evaluation, never as model input.
        missing_targets = [name for name in TARGET_COLUMNS if name not in headers]
        if missing_targets:
            raise ValueError(f"Missing ground truth columns: {missing_targets}")
        target_positions = [headers.index(name) for name in TARGET_COLUMNS]
        targets = pd.DataFrame(
            ([row[index] for index in target_positions]
             for row in sheet.iter_rows(min_row=2, values_only=True)),
            columns=TARGET_COLUMNS,
        ).apply(pd.to_numeric, errors="coerce")
        if not np.isfinite(targets.to_numpy(dtype=float)).all():
            raise ValueError("Every data row must have finite ground truth values")
        noise_error = np.abs(noise - targets["target_total_noise_m"].to_numpy())
        snow_error = np.abs(snow - targets["target_snow_depth_m"].to_numpy())

        last_column = sheet.max_column
        for offset, name in enumerate(NEW_COLUMNS, start=1):
            header = sheet.cell(row=1, column=last_column + offset, value=name)
            header._style = copy(sheet.cell(row=1, column=last_column)._style)
            header.fill = PatternFill(fill_type="solid", fgColor=PREDICTION_HEADER_COLOR)
            font = copy(header.font)
            font.color = "FFFFFFFF"
            font.bold = True
            header.font = font
            sheet.column_dimensions[header.column_letter].width = 27
        for row_number, values in enumerate(zip(noise, snow, noise_error, snow_error), start=2):
            for offset, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_number, column=last_column + offset, value=float(value))
                cell.number_format = "0.000000000"
        # Preserve full-precision predictions, including tiny floating-point residuals.
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(OUTPUT)
        print(f"Saved {len(x):,} predictions: {OUTPUT}")
        print(f"Snow height range (m): {snow.min():.12g} to {snow.max():.12g}")
        print(f"Total noise MAE (m): {noise_error.mean():.12g}")
        print(f"Snow height MAE (m): {snow_error.mean():.12g}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
